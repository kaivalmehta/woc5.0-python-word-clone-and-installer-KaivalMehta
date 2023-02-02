#IMPORTING NECESSARY LIBRARIES
from tkinter import *
from tkinter import filedialog
from tkinter import font
from tkinter import ttk
import PyPDF2
import docx
import re
from fpdf import FPDF
import enchant
#DECLARING LIST OF FONTS AND SIZE
fontlist=[]
sizelist=[2,4,6,8,10,12,14,15,16,18,20,24,28,32,36,40,44,50]

#DECLARING THE NECESSARY VARIABLES
size=20
fontname='Arial'
darkmode=False
spellcheck_button = True

#CREATING THE ROOT WINDOW
root=Tk()
root.title("WOC Word Clone")
root.geometry("1000x1000")
root.config(bg="blue")
#VARIABLES FOR STYLE MENU
bold=IntVar()
italic=IntVar()
underline=IntVar()

#FUNCTIONS USED IN SEVERAL COMMANDS
def newfile():             #CREATING A NEW FILE
    mytext.delete("1.0",END)
    global name
    name=False
    root.title('Untitled')
def openfile():            #OPEN AN EXISTING FILE
    mytext.delete("1.0",END)
    txtfile=filedialog.askopenfilename(initialdir="C:/Downloads/",title="Open File",filetypes=(("Text Files","*.txt"),("HTML Files","*.html")))
    root.title(txtfile)
    global name
    name=txtfile
    txtfile=open(txtfile,'r')
    data=txtfile.read()
    mytext.insert(END,data)
    txtfile.close()
def openpdf():                  #OPENING PDF
    mytext.delete("1.0",END)
    myfile=filedialog.askopenfilename(initialdir="C:/Downloads/",title="Open PDF",filetype=(("PDF    Files","*.pdf"),("All Files","*.*")))
    if myfile:
        pdffile=PyPDF2.PdfReader(myfile)
        text = ""
        for page in range(len(pdffile.pages)):
            text += pdffile.pages[page].extract_text()
        mytext.insert("1.0", text)
def opendocx():                 #OPENING DOCX FILE
    mytext.delete("1.0",END)
    docx_file=filedialog.askopenfilename(filetypes=(("Docx Files", "*.docx"),("All Files","*.*")))
    doc=docx.Document(docx_file)
    text=""
    for data in doc.paragraphs:
        text+=data.text+'\n'
    mytext.insert("1.0",text)   

def saveas():               #SAVING CHANGES AND CREATING A FILE
    txtfile=filedialog.asksaveasfilename(defaultextension=".*",initialdir="C:/Downloads/",title="Save File",filetypes=(("Text Files","*.txt"),("HTML Files","*.html")))
    root.title(txtfile)
    txtfile=open(txtfile,'w')
    txtfile.write(mytext.get(1.0,END))
    txtfile.close()
def savefile():             # SAVING CHANGES IN AN EXISTING FILE OR SAVING AS NEW FILE IF FILE DOESNT EXIST
    global name
    if name :
        txtfile=open(name,'w')
        txtfile.write(mytext.get(1.0,END))
        txtfile.close
    else :
        saveas()
def saveaspdf():            #SAVING AS PDF
    content = mytext.get("1.0", "end-1c")
    file_name = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")])
    if file_name:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, txt=content, ln=1, align="L")
        pdf.output(file_name)
def changefont(fonttype):   # FOR CHANGING FONT TYPE OF THE TEXT
    global fontname
    fontname=fonttype
    mytext.config(font=(fonttype,size))
def changesize(sizeno):   # FOR CHANGING FONT SIZE OF THE TEXT
    global fontname
    global size
    size=sizeno
    mytext.config(font=(fontname,size))
def style():                # FOR CHANGING STYLE OF THE TEXT
    global fontname
    global size
    mytext.tag_configure("bold", font=(fontname, size, "bold"))
    mytext.tag_configure("italic", font=(fontname, size, "italic"))
    mytext.tag_configure("underline", font=(fontname, size, "underline"))
    mytext.tag_configure("bolditalic", font=(fontname, size, "bold italic"))
    mytext.tag_configure("boldunderline", font=(fontname, size, "bold underline"))
    mytext.tag_configure("italicunderline", font=(fontname, size, "italic underline"))
    mytext.tag_configure("bolditalicunderline", font=(fontname,size, "bold italic underline"))
    mytext.tag_remove("bold", "1.0", "end")
    mytext.tag_remove("italic", "1.0", "end")
    mytext.tag_remove("underline", "1.0", "end")
    mytext.tag_remove("bolditalic", "1.0", "end")
    mytext.tag_remove("boldunderline", "1.0", "end")
    mytext.tag_remove("italicunderline", "1.0", "end")
    mytext.tag_remove("bolditalicunderline", "1.0", "end")
    if bold.get() == 1 and italic.get() == 1 and underline.get() == 1:
        mytext.tag_add("bolditalicunderline", "sel.first", "sel.last")
    elif bold.get() == 1 and italic.get() == 1:
        mytext.tag_add("bolditalic", "sel.first", "sel.last")
    elif bold.get() == 1 and underline.get() == 1:
        mytext.tag_add("boldunderline", "sel.first", "sel.last")
    elif italic.get() == 1 and underline.get() == 1:
        mytext.tag_add("italicunderline","sel.first", "sel.last")
    elif bold.get() == 1:
        mytext.tag_add("bold", "sel.first", "sel.last")
    elif italic.get() == 1:
        mytext.tag_add("italic", "sel.first", "sel.last")
    elif underline.get() == 1:
        mytext.tag_add("underline", "sel.first", "sel.last")
    else:
        mytext.config(font=(fontname, size))

def modechange():               #FOR CHANGING THEME/MODE
    global darkmode
    darkmode = not darkmode

    if darkmode:
        mode_button.config(text="Light Mode")
        mytext.config(bg="black", fg="white",insertbackground="white")
    else:
        mode_button.config(text="Dark Mode")
        mytext.config(bg="white", fg="black",insertbackground="black")

def change_color(color):        #FOR CHANGING COLOUR OF THE SELECTED TEXT
    mytext.tag_config("selected", foreground=color)
def text_color():
    colorselect = color.get()
    mytext.tag_add("selected", "sel.first", "sel.last")
    mytext.tag_config("selected", foreground=colorselect)
def highlight_color():          #FOR HIGHLIGHTING SELECTED TEXT
    colorselect = color.get()
    mytext.tag_add("highlight", "sel.first", "sel.last")
    mytext.tag_config("highlight", background=colorselect)
def find_text():                #FOR FINDINF A WORD
    search_term = find_entry.get()
    

    mytext.tag_remove("found", "1.0", "end")

    start_index = "1.0"
    count = 0
    max_count = 100
    while count < max_count:
        match = re.search(search_term, mytext.get(start_index, "end"))
        if not match:
            break
        start_index = f"{start_index} + {match.start()} chars"
        end_index = f"{start_index} + {len(search_term)} chars"
        if "found" not in mytext.tag_names(start_index):
            mytext.tag_add("found", start_index, end_index)
            mytext.tag_configure("found",background="yellow")
        start_index = end_index
        count+=1
    

    
def replace_text():         #FOR REPLACING THE FOUND WORD
    search_term = find_entry.get()
    replace_term = replace_entry.get()


    start_index = "1.0"
    count = 0
    max_count = 100
    while count < max_count:
        match = re.search(search_term, mytext.get(start_index, "end"))
        if not match:
            break
        start_index = f"{start_index} + {match.start()} chars"
        end_index = f"{start_index} + {len(search_term)} chars"
        mytext.delete(start_index, end_index)
        mytext.insert(start_index, replace_term)
        start_index = end_index
        find_entry.delete(0,'end')
        replace_entry.delete(0,'end')
        count += 1
def insertimage():              #FOR INSERTING AN IMAGE
    global my_image
    filepath=filedialog.askopenfilename()
    my_image=PhotoImage(file=filepath)
    position=mytext.index(INSERT)
    mytext.image_create(position,image=my_image)

def spellcheck_button_func():      #FOR SPELL CHECK
    global spellcheck_button
    spellcheck_button = not spellcheck_button
    if spellcheck_button:
        spellcheckbutton.config(text="Spell Check: OFF")
        mytext.tag_remove("underline", "1.0", "end")
    else:
        spellcheckbutton.config(text="Spell Check: ON")
        spellcheck()
def spellcheck():
    txtcont = mytext.get("1.0", "end")
    words = txtcont.split()
    start_index = "1.0"
    for word in words:
        end_index = f"{start_index} + {len(word)}c"
        if not dictionary.check(word):
            mytext.tag_add("underline", start_index, end_index)
        start_index = f"{end_index} + 1c"
def status_bar():           #STATUS BAR DISPLAYING WORDS,LINES AND COLUMNS
    words = len(mytext.get("1.0", "end-1c").split())
    lines = mytext.index("end-1c").split(".")[0]
    columns = mytext.index("insert").split(".")[1]
    status_label.config(text=f"Words: {words} | Lines: {lines} | Columns: {columns}")
    mytext.after(100, status_bar)
# Create a Frame widget
frame = Frame(root)
frame.grid(row=1,column=0,columnspan=9000)
frame.config(width=1540,height=720,padx=5,pady=5)
frame.grid_propagate(False)
# Create a Text widget
mytext=Text(frame,font=("Arial",size),background="white",foreground="black")

mytext.config(width=2000,height=4000,wrap=NONE)
mytext.configure(insertbackground='black')
#Create the vertical scrollbar
v_scrollbar = Scrollbar(frame,width=20, orient="vertical", command=mytext.yview)
mytext.configure(yscrollcommand=v_scrollbar.set)

# Create the horizontal scrollbar
h_scrollbar = Scrollbar(frame, orient="horizontal", command=mytext.xview)
mytext.configure(xscrollcommand=h_scrollbar.set)


mytext.grid(row=0, column=0,sticky="nsew")
v_scrollbar.grid(row=0, column=1,sticky="ns")
h_scrollbar.grid(row=2, column=0,sticky="ew")


color=StringVar()

menubuttons_frame = Frame(root,background="blue")
menubuttons_frame.grid(row=0, column=0)


status_label = Label(frame, text="", bd=1, relief="sunken", anchor="w")
status_label.grid(row=1,column=0)
 
mytext.config(wrap="word")
mytext.tag_configure("underline", underline=True, foreground="red")
dictionary = enchant.Dict("en_US")

find_entry = Entry(menubuttons_frame)
find_entry.grid(row=0,column=19)
find_button = Button(menubuttons_frame, text="Find", command=find_text,font=("Arial",15),background="blue",foreground="white",padx=17)
find_button.grid(row=0,column=18,padx=25)

replace_entry = Entry(menubuttons_frame)
replace_entry.grid(row=1,column=19)

replace_button = Button(menubuttons_frame, text="Replace", command=replace_text,font=("Arial",15),background="blue",foreground="white")
replace_button.grid(row=1,column=18,padx=25)

insert_image_button = Button(menubuttons_frame, text="Insert Image", command=insertimage,font=("Arial",15),background="blue",foreground="white")
insert_image_button.grid(row=0,column=17,padx=25)
spellcheckbutton=Button(menubuttons_frame,text="Spell Check: OFF ",command=spellcheck_button_func,font=("Arial",15),background="blue",foreground="white")
spellcheckbutton.grid(row=0,column=10,padx=30)
frame.grid_columnconfigure(0, weight=1)
frame.grid_rowconfigure(0, weight=1)
colormenu=ttk.OptionMenu(menubuttons_frame, color, "red", "green", "blue", "purple","black","orange","yellow","white", command=change_color)
colormenu.grid(row=1,column=5)
highlight_button = Button(menubuttons_frame, text="Highlight", command=highlight_color,font=("Arial",15),background="blue",foreground="white")
highlight_button.grid(row=0,column=4)
colourtext_button = Button(menubuttons_frame, text="Change Colour", command=text_color,font=("Arial",15),background="blue",foreground="white")
colourtext_button.grid(row=0,column=6)

file_menubutton = Menubutton(menubuttons_frame, text="File   ", font=("Arial", 15),background="blue",foreground="white")
file_menubutton.grid(row=0,column=0)



file_menu = Menu(file_menubutton, tearoff=0)
file_menubutton.config(menu=file_menu)

# Add options to the File menu
file_menu.add_command(label="New", command=newfile)
file_menu.add_command(label="Open", command=openfile)
file_menu.add_command(label="Open PDF", command=openpdf)
file_menu.add_command(label="Open Docx", command=opendocx)
file_menu.add_command(label="Save", command=savefile)
file_menu.add_command(label="Save As", command=saveas)
file_menu.add_command(label="Save As PDF", command=saveaspdf)
file_menu.add_separator()
file_menu.add_command(label="Exit", command=root.destroy)


font_menubutton = Menubutton(menubuttons_frame, text="Font  ", font=("Arial", 15),background="blue",foreground="white")
font_menubutton.grid(row=0,column=1)


font_menu = Menu(font_menubutton, tearoff=0)
font_menubutton.config(menu=font_menu)

# Add options to the Font menu
for i in font.families():
    fontlist.append(i)

for font in fontlist:
        font_menu.add_command(label=font, command=lambda fonttype=font: changefont(fonttype))



size_menubutton = Menubutton(menubuttons_frame, text="Size   ", font=("Arial", 15),background="blue",foreground="white")
size_menubutton.grid(row=0,column=2)


size_menu = Menu(size_menubutton, tearoff=0)
size_menubutton.config(menu=size_menu)


for sizeno in sizelist:
    size_menu.add_command(label=sizeno, command=lambda sizeno=sizeno: changesize(sizeno))


style_menubutton = Menubutton(menubuttons_frame, text="Style  ", font=("Arial", 15),background="blue",foreground="white")
style_menubutton.grid(row=0,column=3)


style_menu = Menu(style_menubutton, tearoff=0)
style_menubutton.config(menu=style_menu)
mode_button=Button(menubuttons_frame,text="Dark Mode",font=("Arial", 15),background="blue",foreground="white",command=modechange)
mode_button.grid(row=0,column=30,padx=30)
# Add options to the Style menu
style_menu.add_checkbutton(label="Bold", variable=bold, command=style)
style_menu.add_checkbutton(label="Italic", variable=italic, command=style)
style_menu.add_checkbutton(label="Underline", variable=underline, command=style)
mytext.bind("<<Key>>", status_bar)
mytext.bind("<Button-1>", status_bar)
mytext.after(100,status_bar)
root.mainloop()